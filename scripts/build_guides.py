"""Build illustrated Word manuals from docs/admin.md and docs/user.md.
Requires python-docx. Run from any directory; render resulting DOCX for layout QA.
"""
from pathlib import Path
import re
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT = Path(__file__).resolve().parents[1]

def plain(s):
    return re.sub(r'\[([^]]+)\]\([^)]+\)', r'\1', s).replace('`', '')

def build(source, target):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(1.65)
    sec.left_margin = sec.right_margin = Cm(1.905)
    for name, size in [('Normal',10.5),('Title',26),('Heading 1',19),('Heading 2',12),('Caption',9)]:
        st=doc.styles[name];st.font.name='Arial';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
        st.paragraph_format.space_after=Pt(7)
    doc.styles['Normal'].paragraph_format.line_spacing=1.08
    footer=sec.footer.paragraphs[0];footer.alignment=2
    r=footer.add_run('sambaadm  |  ');r.font.size=Pt(8)
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
    for el in doc.styles.element.iter():
        for child in list(el):
            if child.tag == qn('w:pBdr'): el.remove(child)
    headings=[line[3:] for line in source.read_text().splitlines() if line.startswith('## ')]
    code=False
    contents=False
    for line in source.read_text().splitlines():
        if not line: continue
        if line.startswith('```'): code=not code;continue
        if line.startswith('# '): doc.add_paragraph(line[2:],'Title');continue
        if line.startswith('## '):
            if not contents:
                doc.add_paragraph('Содержание', 'Heading 2')
                for i, title in enumerate(headings, 2): doc.add_paragraph(f'{i}  {title}')
                contents=True
            p=doc.add_paragraph(line[3:],'Heading 1');p.paragraph_format.page_break_before=True;continue
        m=re.match(r'!\[(.*?)\]\((.*?)\)',line)
        if m:
            p=doc.add_paragraph();p.paragraph_format.keep_with_next=True
            if "share-wizard" in m[2]: p.alignment=1
            p.add_run().add_picture(str(source.parent/m[2]),width=Cm(13.5 if 'share-wizard' in m[2] else 17.0))
            doc.add_paragraph(m[1], 'Caption');continue
        if line.startswith('[Иллюстрированная') or line.startswith('[Инструкция Word'): continue
        p=doc.add_paragraph(plain(line))
        if code is False and '](' in line:
            for label, href in re.findall(r'\[([^]]+)\]\(([^)]+)\)',line):
                url=href if href.startswith('http') else 'https://github.com/rkislov/sambaadm/blob/main/'+str((source.parent/href).resolve().relative_to(ROOT))
                r=p.add_run(' '+url);r.font.size=Pt(8)
        if code:
            p.paragraph_format.space_after=Pt(0)
            p.paragraph_format.line_spacing=1
            for r in p.runs:r.font.name='Courier New';r.font.size=Pt(8)
    doc.core_properties.title=doc.paragraphs[0].text
    doc.core_properties.author='sambaadm'
    doc.core_properties.subject='Руководство для main и v0.23.0-dev от 21 октября 2026 года'
    target.parent.mkdir(parents=True,exist_ok=True);doc.save(target)

build(ROOT/'docs/admin.md',ROOT/'docs/admin-guide/sambaadm-admin-guide.docx')
build(ROOT/'docs/user.md',ROOT/'docs/user-guide/sambaadm-user-guide.docx')
